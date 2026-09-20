"""Export the elementary knowledge graph to Excel, with 지식·이해/과정·기능/가치·태도
taken from data/curriculum-content-systems (grade-band split, elementary) instead of
the graph's own lists (which copy the 초·중 공통 group, i.e. mostly 중학교 items).

Sheets: 안내 / 노드 / 엣지 / 성취기준 / 핵심아이디어 / 내용체계원본 / 검증 / 교과
"""
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
GRAPH = ROOT / "data" / "elementary_knowledge_graph.json"
CS_DIR = ROOT / "data" / "curriculum-content-systems"
CUR_DIR = ROOT / "public" / "curriculum_json"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "exports" / "초등_지식그래프_성취기준_노드엣지.xlsx"

g = json.loads(GRAPH.read_text(encoding="utf-8"))
meta = g["metadata"]
subjects = {s["id"]: s for s in g["subjects"]}
nodes = {n["id"]: n for n in g["nodes"]}
standards = {a["id"]: a for a in g["achievementStandards"]}
core_ideas = {c["id"]: c for c in g["coreIdeas"]}
cross_links = {l["id"]: l for l in g["links_cross_subject"]}

# curriculum_json groups (source of the graph's functions/attitudes lists) — used only for the 검증 sheet
cur_groups = {}
for f in CUR_DIR.glob("*.json"):
    for grp in json.loads(f.read_text(encoding="utf-8")).get("core_idea_groups", []):
        cur_groups[grp["id"]] = grp

# ---------------------------------------------------------------- content systems (ground truth)
CAT_KEYS = ("지식⋅이해", "과정⋅기능", "가치⋅태도")
GRAPH_SUBJECT_NAME = {
    "sub_kor": "국어", "sub_math": "수학", "sub_sci": "과학", "sub_soc": "사회", "sub_mor": "도덕",
    "sub_art": "미술", "sub_mus": "음악", "sub_pe": "체육", "sub_eng": "영어", "sub_prac": "실과",
    "sub_int": "통합교과", "sub_extra": "창의적 체험활동",
}
# graph area → content-system area(s). Anything not listed falls back to normalized equality/containment.
AREA_ALIAS = {
    ("사회", "자연·인문환경과 인간생활"): ["자연환경과 인간생활", "인문환경과 인간생활"],
    ("사회", "인문환경·지속가능한 세계"): ["인문환경과 인간생활", "지속가능한 세계"],
    ("영어", "이해"): ["이해(reception)"],
    ("영어", "표현"): ["표현(production)"],
}
BAND_OF_GRADE = {"초1-2": "1-2학년군", "초3-4": "3-4학년군", "초5-6": "5-6학년군"}


def nk(s):
    return re.sub(r"[\s⋅·•・()（）\[\]]", "", s or "").lower()


cs_records = []  # elementary records with bands
for f in sorted(CS_DIR.glob("*.json")):
    d = json.loads(f.read_text(encoding="utf-8"))
    subj = d["메타"]["교과"]
    for i, r in enumerate(d["내용체계"]):
        bands = r.get("학년군별") or {}
        if not bands or str(r.get("교육과정", "")).startswith("선택 중심 교육과정"):
            continue
        cs_records.append({
            "id": f"{f.name}#{i}", "file": f.name, "subject": subj, "course": r.get("과목", ""),
            "area": r.get("영역", ""), "pages": r.get("출처쪽", []), "core_ideas": r.get("핵심아이디어", []),
            "bands": {b: {k: v.get(k, []) for k in CAT_KEYS} for b, v in bands.items() if "중" not in b},
        })


def cs_for(subject_id, area):
    """Content-system records for a graph (subject, area). Returns list (may be several, e.g. 통합교과 3 과목)."""
    subj = GRAPH_SUBJECT_NAME.get(subject_id, subject_id)
    wanted = AREA_ALIAS.get((subj, area), [area])
    out = []
    for w in wanted:
        nw = nk(w)
        hits = [r for r in cs_records if r["subject"] == subj and nk(r["area"]) == nw]
        if not hits:
            hits = [r for r in cs_records if r["subject"] == subj and (nw in nk(r["area"]) or nk(r["area"]) in nw)]
        out.extend(h for h in hits if h not in out)
    return out


def band_items(records, band, cat):
    items = []
    for r in records:
        for it in r["bands"].get(band, {}).get(cat, []):
            if it not in items:
                items.append(it)
    return items


# ---------------------------------------------------------------- styling helpers
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
BODY_FONT = Font(name="Arial", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
NODE_TYPE_KO = {"subject": "교과", "core_idea": "핵심아이디어", "standard": "성취기준"}
ROLE_KO = {"phenomenon": "현상(탐구 대상)", "tool": "도구(방법·수단)", "expression": "표현(산출·표현)", "value": "가치(판단·태도)", "integrated": "통합"}
REL_KO = {"has_core_idea": "교과 → 핵심아이디어 (소속)", "has_standard": "핵심아이디어 → 성취기준 (소속)"}
SEM_KO = {"semantic_similarity": "의미 유사(임베딩)", "shares_keyword": "키워드 공유"}
METHOD_KO = {"explicit": "명시적(교육과정 구조)", "hybrid_embedding70_tfidf30": "하이브리드 유사도(임베딩 70% + TF-IDF 30%)"}


def join(lst):
    return "; ".join(lst) if lst else ""


def subj_name(sid):
    return subjects.get(sid, {}).get("name_ko", sid)


def node_label(nid):
    if nid is None:
        return "(핵심아이디어 미배정)"
    n = nodes.get(nid)
    if n is None:
        return nid
    if n["type"] == "standard":
        return standards.get(nid, {}).get("code", n["label"])
    return n["label"]


def node_text(nid):
    if nid is None:
        return "원본 JSON에서 source가 null (창의적 체험활동 성취기준은 핵심아이디어에 연결되지 않음)"
    n = nodes.get(nid)
    if n is None:
        return ""
    if n["type"] == "standard":
        return standards.get(nid, {}).get("text", n.get("text", ""))
    return n["label"]


def write_sheet(ws, headers, rows, widths, wrap_cols=(), freeze="A2"):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    wrap_idx = set(wrap_cols)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = WRAP if cell.column in wrap_idx else TOP
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()

# ---------------------------------------------------------------- 안내
ws = wb.active
ws.title = "안내"
n_expl = sum(1 for e in g["edges"] if e["method"] == "explicit")
info = [
    ["초등 융합 지식 그래프 — 성취기준 노드·엣지 내보내기 (내용체계 학년군 보정판)"],
    [],
    ["■ 데이터 출처"],
    ["항목", "출처 파일", "비고"],
    ["노드·엣지·성취기준 문장·핵심아이디어 문장", "data/elementary_knowledge_graph.json", f"버전 {meta['version']}"],
    ["지식·이해 / 과정·기능 / 가치·태도", "data/curriculum-content-systems/*내용체계.json", "2022 개정 교육과정 별책 PDF의 내용 체계 표에서 학년군별로 추출한 자료. 앱의 [내용체계] 컨텍스트와 A-2-1 확정 매핑이 쓰는 것과 같은 파일"],
    [],
    ["■ 지식·이해 / 과정·기능 / 가치·태도의 출처"],
    ["그래프 JSON의 coreIdeas[].knowledge 와 achievementStandards[].knowledge/functions/competencies 는 원래 public/curriculum_json 의 핵심아이디어 그룹 값을 그대로 복사한 것이었습니다. 그 그룹은 초·중 공통 교육과정 단위라 학년군 구분이 없고 실제 내용은 대부분 중학교 1~3학년 열입니다(예: 과학 '옴의 법칙', 수학 '소인수분해'). 그래서 이 엑셀은 처음부터 내용체계 파일(학년군별)에서 항목을 가져오며, 그래프 JSON 자체도 같은 기준으로 정화되어 있습니다(metadata.content_lists_sanitized_at). 상위 원본 그룹 항목의 판정은 '검증' 시트에 있습니다."],
    ["검증 루프", "npm run verify:curriculum — 데이터 파일과 실제 프롬프트 생성 경로에 중학교·학년군 밖 항목이 없는지 검사. 실패 시 빌드가 중단됨."],
    [],
    ["■ 규모"],
    ["교과 수", meta["subjects_count"]], ["성취기준 수", meta["standards_count"]], ["노드 수", len(g["nodes"])], ["엣지 수", len(g["edges"])],
    ["  - 구조 엣지(explicit)", n_expl], ["  - 교과간 연결 엣지", meta["cross_subject_edges"]],
    ["유사도 방법", meta["similarity_method"]], ["임베딩 가중치(alpha)", meta["embedding_alpha"]],
    ["동일 학년군 임계값", meta["same_grade_threshold"]], ["교차 학년군 임계값", round(meta["cross_grade_threshold"], 3)],
    [],
    ["■ 노드(Node) 유형"],
    ["유형", "설명", "개수"],
    ["교과(subject)", "12개 교과. 역할(현상·도구·표현·가치·통합)로 구분", sum(1 for n in g["nodes"] if n["type"] == "subject")],
    ["핵심아이디어(core_idea)", "교과 내 영역별 핵심 아이디어 묶음", sum(1 for n in g["nodes"] if n["type"] == "core_idea")],
    ["성취기준(standard)", "2022 개정 교육과정 초등 성취기준", sum(1 for n in g["nodes"] if n["type"] == "standard")],
    [],
    ["■ 엣지(Edge) 관계 유형"],
    ["관계", "설명", "개수", "생성 방법"],
    ["has_core_idea", "교과 → 핵심아이디어 소속 관계", sum(1 for e in g["edges"] if e["relation"] == "has_core_idea"), "명시적"],
    ["has_standard", "핵심아이디어 → 성취기준 소속 관계", sum(1 for e in g["edges"] if e["relation"] == "has_standard"), "명시적"],
]
for rel, desc in meta["relation_types"].items():
    info.append([rel, desc, sum(1 for e in g["edges"] if e["relation"] == rel), "하이브리드 유사도"])
info += [
    [],
    ["■ 시트 구성"],
    ["시트", "내용"],
    ["노드", "그래프의 모든 노드(교과·핵심아이디어·성취기준)"],
    ["엣지", "그래프의 모든 엣지. 출발·도착 노드의 라벨·문장, 가중치, 근거(공유 키워드)"],
    ["성취기준", "성취기준 627개. 지식·이해/과정·기능/가치·태도는 해당 성취기준의 교과·영역·학년군에 해당하는 내용 체계 항목(영역×학년군 단위이며 성취기준 1개 단위가 아님)"],
    ["핵심아이디어", "핵심아이디어 54개 × 학년군. 학년군별 지식·이해/과정·기능/가치·태도"],
    ["내용체계원본", "curriculum-content-systems 파일의 초등 기록 전체(교과·과목·영역·학년군·출처쪽)"],
    ["검증", "상위 원본(public/curriculum_json 핵심아이디어 그룹) 목록의 항목별 판정: 초등 내용체계와 일치 / 중학교 / 초등 내용체계에 없음. 그래프 JSON은 이미 정화됐으므로 이 시트는 원본 오염의 기록"],
    ["교과", "교과 12개와 역할"],
    [],
    ["※ 가중치(weight)는 0~1 범위의 유사도 점수. 구조 엣지는 1.0 고정."],
    ["※ 여러 값을 가진 칸은 세미콜론(;)으로 구분."],
    ["※ 데이터 유의 1: 창의적 체험활동 성취기준 18개는 원본 JSON에서 핵심아이디어가 배정되지 않아 has_standard 엣지의 출발 노드가 null임. 엣지 시트에는 '(null)'로 표기. 내용 체계도 학년군 구분이 없어 비어 있음."],
    ["※ 데이터 유의 2: 사회의 그래프 영역 '자연·인문환경과 인간생활', '인문환경·지속가능한 세계'는 내용 체계 영역 두 개를 합친 이름이라 두 영역의 항목을 함께 실었음. 앱의 매칭 코드는 이름 포함 여부로 한 영역만 잡음."],
]
for row in info:
    ws.append(row)
for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
    for cell in row:
        cell.font, cell.alignment = BODY_FONT, WRAP
ws["A1"].font = Font(name="Arial", bold=True, size=14)
for r in range(1, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value
    if isinstance(v, str) and v.startswith("■"):
        ws.cell(row=r, column=1).font = Font(name="Arial", bold=True, size=11)
        nxt = ws.cell(row=r + 1, column=1).value
        if isinstance(nxt, str) and nxt in ("항목", "유형", "관계", "시트"):
            for c in range(1, 5):
                hc = ws.cell(row=r + 1, column=c)
                if hc.value:
                    hc.font, hc.fill = HEADER_FONT, HEADER_FILL
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 60
ws.column_dimensions["C"].width = 60
ws.column_dimensions["D"].width = 22

# ---------------------------------------------------------------- 노드
ws = wb.create_sheet("노드")
rows = []
for n in g["nodes"]:
    t = n["type"]
    a = standards.get(n["id"], {})
    rows.append([
        n["id"], NODE_TYPE_KO.get(t, t), node_label(n["id"]),
        n["label"] if t == "subject" else subj_name(n.get("subject_id", "")),
        ROLE_KO.get(n.get("role", ""), n.get("role", "")) if t == "subject" else "",
        n.get("grade_band", ""), n.get("area", ""),
        a.get("text", "") if t == "standard" else "",
        join(n.get("keywords", [])),
    ])
write_sheet(ws, ["노드 ID", "노드 유형", "라벨(코드/이름)", "교과", "교과 역할", "학년군", "영역", "성취기준 문장", "키워드"],
            rows, [26, 14, 24, 16, 16, 9, 16, 60, 36], wrap_cols=(8, 9))

# ---------------------------------------------------------------- 엣지
ws = wb.create_sheet("엣지")
rows = []
for e in g["edges"]:
    s, t = e["source"], e["target"]
    ev = cross_links.get(e["id"], {}).get("evidence", {})
    rows.append([
        e["id"], REL_KO.get(e["relation"], e["relation"]),
        s if s is not None else "(null)", NODE_TYPE_KO.get(nodes.get(s, {}).get("type", ""), ""), node_label(s),
        subj_name(nodes.get(s, {}).get("subject_id", s)), node_text(s),
        t if t is not None else "(null)", NODE_TYPE_KO.get(nodes.get(t, {}).get("type", ""), ""), node_label(t),
        subj_name(nodes.get(t, {}).get("subject_id", t)), node_text(t),
        e.get("weight"), METHOD_KO.get(e.get("method", ""), e.get("method", "")),
        SEM_KO.get(e.get("relation_sem", ""), e.get("relation_sem", "")), e.get("grade_band", ""),
        ("예" if e.get("same_grade_band") else "아니오") if "same_grade_band" in e else "",
        join(ev.get("shared_keywords", [])), join(ev.get("shared_functions", [])), join(ev.get("shared_knowledge", [])),
    ])
write_sheet(ws, ["엣지 ID", "관계 유형", "출발 노드 ID", "출발 유형", "출발 라벨", "출발 교과", "출발 문장",
                 "도착 노드 ID", "도착 유형", "도착 라벨", "도착 교과", "도착 문장",
                 "가중치", "생성 방법", "의미 관계", "학년군", "동일 학년군", "공유 키워드", "공유 기능", "공유 지식"],
            rows, [14, 24, 26, 12, 22, 14, 48, 26, 12, 22, 14, 48, 9, 26, 16, 9, 10, 28, 28, 28],
            wrap_cols=(7, 12, 18, 19, 20), freeze="C2")
for row in ws.iter_rows(min_row=2, min_col=13, max_col=13):
    for cell in row:
        cell.number_format = "0.0000"

# ---------------------------------------------------------------- 성취기준 (content-system by band)
ws = wb.create_sheet("성취기준")
rows = []
unmatched = []
for a in g["achievementStandards"]:
    ci = core_ideas.get(a.get("core_idea_id"), {})
    recs = cs_for(a["subject_id"], a.get("area", ""))
    band = BAND_OF_GRADE.get(a.get("grade_band", ""), "")
    kn = band_items(recs, band, "지식⋅이해")
    fn = band_items(recs, band, "과정⋅기능")
    at = band_items(recs, band, "가치⋅태도")
    src = "; ".join(f"{r['course']}/{r['area']}" for r in recs)
    if not recs and a["subject_id"] != "sub_extra":
        unmatched.append((a["code"], subj_name(a["subject_id"]), a.get("area")))
    rows.append([
        a["id"], a["code"], subj_name(a["subject_id"]), a.get("grade_band", ""), a.get("area", ""),
        a.get("core_idea_id", ""), join(ci.get("ideas", [])), a.get("text", ""), join(a.get("keywords", [])),
        band, join(kn), join(fn), join(at), src,
    ])
write_sheet(ws, ["노드 ID", "성취기준 코드", "교과", "학년군", "영역", "핵심아이디어 ID", "핵심아이디어 문장", "성취기준 문장", "키워드",
                 "내용체계 학년군", "지식·이해(영역×학년군)", "과정·기능(영역×학년군)", "가치·태도(영역×학년군)", "내용체계 출처(과목/영역)"],
            rows, [26, 14, 12, 8, 18, 24, 50, 60, 30, 11, 50, 50, 44, 26],
            wrap_cols=(7, 8, 9, 11, 12, 13, 14), freeze="C2")

# ---------------------------------------------------------------- 핵심아이디어 (× band)
ws = wb.create_sheet("핵심아이디어")
std_count = defaultdict(int)
for a in g["achievementStandards"]:
    std_count[(a.get("core_idea_id"), a.get("grade_band"))] += 1
rows = []
for c in g["coreIdeas"]:
    recs = cs_for(c["subject_id"], c.get("area", ""))
    bands = sorted({b for r in recs for b in r["bands"]}) or [""]
    for b in bands:
        grade = next((k for k, v in BAND_OF_GRADE.items() if v == b), "")
        if b and std_count.get((c["id"], grade), 0) == 0 and not band_items(recs, b, "지식⋅이해"):
            continue  # 해당 학년군에 성취기준도 지식·이해도 없음 (예: 과학 1-2학년군)
        rows.append([
            c["id"], subj_name(c["subject_id"]), c.get("area", ""), b, std_count.get((c["id"], grade), 0),
            join(c.get("ideas", [])),
            join(band_items(recs, b, "지식⋅이해")), join(band_items(recs, b, "과정⋅기능")), join(band_items(recs, b, "가치⋅태도")),
            "; ".join(f"{r['course']}/{r['area']}" for r in recs),
        ])
write_sheet(ws, ["노드 ID", "교과", "영역", "학년군", "해당 학년군 성취기준 수", "핵심아이디어 문장", "지식·이해", "과정·기능", "가치·태도", "내용체계 출처(과목/영역)"],
            rows, [30, 12, 20, 11, 10, 60, 50, 50, 44, 26], wrap_cols=(6, 7, 8, 9, 10), freeze="D2")

# ---------------------------------------------------------------- 내용체계원본
ws = wb.create_sheet("내용체계원본")
rows = []
for r in cs_records:
    for b, cats in r["bands"].items():
        rows.append([r["file"], r["subject"], r["course"], r["area"], b,
                     join(cats["지식⋅이해"]), join(cats["과정⋅기능"]), join(cats["가치⋅태도"]),
                     ", ".join(str(p) for p in r["pages"]), join(r["core_ideas"])])
write_sheet(ws, ["파일", "교과", "과목", "영역", "학년군", "지식·이해", "과정·기능", "가치·태도", "출처 쪽", "핵심아이디어(원문)"],
            rows, [20, 10, 14, 24, 11, 50, 50, 44, 12, 60], wrap_cols=(6, 7, 8, 10), freeze="E2")

# ---------------------------------------------------------------- 검증 (graph original lists vs content systems)
ws = wb.create_sheet("검증")
elem_pool = defaultdict(lambda: defaultdict(set))
mid_pool = defaultdict(lambda: defaultdict(set))
for f in sorted(CS_DIR.glob("*.json")):
    d = json.loads(f.read_text(encoding="utf-8"))
    subj = d["메타"]["교과"]
    for r in d["내용체계"]:
        for b, cats in (r.get("학년군별") or {}).items():
            pool = mid_pool if "중" in b else elem_pool
            for k in CAT_KEYS:
                for it in cats.get(k, []):
                    pool[subj][k].add(nk(it))


def verdict(item, subj, cat):
    n = nk(item)
    if n in elem_pool[subj][cat]:
        return "초등 내용체계와 일치"
    if any((n in p or p in n) and len(min(n, p, key=len)) >= 3 for p in elem_pool[subj][cat]):
        return "초등 내용체계와 부분 일치"
    if n in mid_pool[subj][cat] or any((n in p or p in n) and len(min(n, p, key=len)) >= 3 for p in mid_pool[subj][cat]):
        return "중학교 내용체계와 일치"
    return "초등 내용체계에 없음"


CAT_OF = {"knowledge": "지식⋅이해", "functions": "과정⋅기능", "attitudes": "가치⋅태도"}
rows = []
for c in g["coreIdeas"]:
    subj = GRAPH_SUBJECT_NAME.get(c["subject_id"], c["subject_id"])
    grp = cur_groups.get(c["id"], {})
    lists = {"knowledge": grp.get("knowledge", []), "functions": grp.get("functions", []), "attitudes": grp.get("attitudes", [])}
    for field, items in lists.items():
        for it in items:
            rows.append([subj, c["id"], c.get("area", ""), CAT_OF[field].replace("⋅", "·"), it, verdict(it, subj, CAT_OF[field]),
                         ", ".join(grp.get("school_levels", []))])
write_sheet(ws, ["교과", "핵심아이디어 ID", "영역", "범주", "상위 원본(curriculum_json 그룹) 항목", "판정", "원본 그룹의 학교급 범위"],
            rows, [10, 30, 20, 10, 50, 22, 18], wrap_cols=(5,), freeze="A2")
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    v = row[5].value
    if v == "중학교 내용체계와 일치":
        row[5].fill = PatternFill("solid", fgColor="F8CBAD")
    elif v == "초등 내용체계에 없음":
        row[5].fill = PatternFill("solid", fgColor="FFE699")
    elif v and v.startswith("초등"):
        row[5].fill = PatternFill("solid", fgColor="C6E0B4")

# ---------------------------------------------------------------- 교과
ws = wb.create_sheet("교과")
sub_std, sub_ci = defaultdict(int), defaultdict(int)
for a in g["achievementStandards"]:
    sub_std[a["subject_id"]] += 1
for c in g["coreIdeas"]:
    sub_ci[c["subject_id"]] += 1
rows = [[s["id"], s["name_ko"], ROLE_KO.get(s.get("role", ""), s.get("role", "")), sub_ci[s["id"]], sub_std[s["id"]]] for s in g["subjects"]]
write_sheet(ws, ["노드 ID", "교과명", "역할", "핵심아이디어 수", "성취기준 수"], rows, [14, 34, 18, 14, 12])

wb.save(OUT)
print(f"saved {OUT}")
print("sheets:", wb.sheetnames)
print("standards without content-system match (excluding 창체):", len(unmatched), unmatched[:10])
